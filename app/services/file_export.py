"""Answers delivered as files: PDF to read, Word to edit, Excel for numbers.

The model writes content; this module makes the file. A model asked to emit a file
directly produces broken spreadsheets and PDFs whose Arabic falls apart — so it is only
ever asked for what it is good at, text and Markdown tables, and everything about the
file itself is done here, the same way every time: right-to-left, Arabic shaping, a
header with the title and date, the sources listed at the end.

Every number placed in a spreadsheet cell is looked for in the passages the answer was
built from. One that is not found is highlighted and noted, so a figure the model
produced without support never sits in a clean cell looking like a quotation.
"""

from __future__ import annotations

import json
import re
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from pathlib import Path

from app.exceptions import DocumentNotFoundError

FORMATS = {"pdf": "PDF", "docx": "Word", "xlsx": "Excel"}
MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
#: Files are kept this long, then removed by the nightly job.
KEEP_DAYS = 7

# -- what was asked for ----------------------------------------------------

_MAKE = (
    r"(?:اعمل|اعملي|إعمل|اعملى|سوّي|سوي|سويلي|هات|هاتلي|هاتلى|ابعت|ابعتلي|ارسل|أرسل|"
    r"حول|حوّل|حولي|حوّلي|صدّر|صدر|طلّع|طلع|طلعلي|جهز|جهّز|جهزلي|انشئ|أنشئ|اكتب|اكتبلي|"
    r"اعطني|أعطني|عايز|عاوز|ابغى|أبغى|ابي|أبي|نزّل|نزل|"
    r"generate|create|make|export|give\s+me|convert|prepare|send\s+me)"
)
_PDF = r"(?:pdf|بي\s*دي\s*[اإ]ف|بى\s*دى\s*[اإ]ف|بدف)"
_WORD = r"(?:word|docx|وورد|ورد)"
_EXCEL = r"(?:excel|xlsx|اكسل|إكسل|اكسيل|إكسيل|اكسال|شيت|spreadsheet)"
_FORMAT = re.compile(rf"(?<![\w.])({_PDF}|{_WORD}|{_EXCEL})(?![\w])", re.IGNORECASE)
_ASKS = re.compile(rf"(?:^|\s|و){_MAKE}", re.IGNORECASE)
_AS_FILE = re.compile(r"(?:بصيغة|على\s+شكل|في\s+ملف|كملف|ملف|as\s+an?|in|to)\s*$", re.IGNORECASE)

_CONVERSATION = re.compile(r"(?:المحادث[ةه]|الشات|المحاور[ةه]|كل\s+(?:الكلام|اللي\s+فات)|conversation|chat)", re.IGNORECASE)
_SUMMARY = re.compile(r"(?:ملخص|لخص|تلخيص|ملخّص|summary|summari[sz]e)", re.IGNORECASE)
_THAT = re.compile(
    r"^(?:دا|ده|دي|هذا|هذه|هذي|الإجابة|الاجابة|الرد|الجواب|الكلام\s+(?:دا|ده|هذا)|الجدول\s+(?:دا|ده|هذا)|"
    r"it|this|that|the\s+answer)?$",
    re.IGNORECASE,
)
#: Whole words dropped from a file request to leave the question it carries.
_FILLER_WORDS = {
    "لي", "لى", "ملف", "ملفات", "بصيغة", "كملف", "فايل", "نسخة", "فيه", "فيها", "عليه", "عليها",
    "please", "a", "an", "the", "file", "as", "in", "to", "into", "for", "me", "of", "with",
    "لـ", "ل", "ب", "بـ", "يا", "من", "فضلك", "لو", "سمحت", "شكل", "على", "في", "صيغة",
}
_FILLER_PHRASES = re.compile(r"(?:على|في)\s+شكل|من\s+فضلك|لو\s+سمحت|بصيغة", re.IGNORECASE)
_VERB = re.compile(rf"^(?:و)?{_MAKE}(?:لي|لى|ي|ني)?$", re.IGNORECASE)


def _without_request_words(text: str) -> str:
    text = _FILLER_PHRASES.sub(" ", text)
    for _ in range(3):
        text = re.sub(rf"^\s*(?:و)?{_MAKE}(?:لي|لى|ي|ني)?(?=\s|$)", "", text, flags=re.IGNORECASE)
    words = text.split()
    # Request words lead the sentence ("اعمل لي ملف …") or trail it ("… كملف"); the
    # same words in the middle belong to the question ("في أمر التغيير").
    while words and (words[0].lower() in _FILLER_WORDS or _VERB.match(words[0])):
        words.pop(0)
    while words and (words[-1].lower() in _FILLER_WORDS or _VERB.match(words[-1])):
        words.pop()
    # "بأسعار …", "بالغرامات": the attached "with" of the request, not part of the word.
    if words and len(words[0]) > 4 and (
        re.match(r"^ب(?=ال|أ|ا)", words[0])
        or words[0][1:] in {"جدول", "قائمة", "ملخص", "مقارنة", "تفاصيل", "بنود", "كل", "جميع", "ارقام", "أرقام"}
    ) and words[0].startswith("ب"):
        words[0] = words[0][1:]
    return " ".join(words)


@dataclass
class FileRequest:
    format: str
    #: "last" — the answer before this one; "conversation" — the whole conversation;
    #: "question" — a new question, answered and delivered as the file.
    subject: str
    question: str = ""


def detect(question: str) -> FileRequest | None:
    """A request for a file, or None for an ordinary question.

    Needs both a verb of making or giving and a format named as such: "ما في ملف
    العقد.pdf؟" names a PDF and asks nothing of the kind, and is answered as before.
    """
    text = " ".join(question.split())
    found = _FORMAT.search(text)
    if found is None or not _ASKS.search(text):
        return None
    word = found.group(1).lower()
    fmt = "pdf" if re.fullmatch(_PDF, word, re.IGNORECASE) else "docx" if re.fullmatch(_WORD, word, re.IGNORECASE) else "xlsx"
    before, after = text[: found.start()], text[found.end():]
    # "لـ PDF", "بال pdf": the joining letters before the format are not the question's.
    before = re.sub(r"(?:\s|^)(?:لـ|ل|بـ|ب|بال|كـ|ك|الى|إلى|to|as|in)\s*$", " ", before, flags=re.IGNORECASE)
    rest = " ".join((before + " " + after).replace("؟", " ").replace("?", " ").split()).strip(" .،,:-—")
    rest = _without_request_words(rest).strip(" .،,:-—")
    if _CONVERSATION.search(rest) or (_SUMMARY.search(rest) and len(rest.split()) <= 3):
        return FileRequest(fmt, "conversation")
    if _THAT.match(rest) or len(rest.split()) <= 1:
        return FileRequest(fmt, "last")
    return FileRequest(fmt, "question", rest + "؟")


# -- what goes in it -------------------------------------------------------

@dataclass
class Block:
    kind: str  # heading | paragraph | bullet | table
    text: str = ""
    level: int = 1
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class Content:
    title: str
    blocks: list[Block]
    sources: list[str] = field(default_factory=list)
    #: The text of the passages the content rests on — what numbers are checked against.
    evidence: str = ""
    author: str = ""
    note: str = ""


_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_CITATION = re.compile(r"\s*\[(\d+(?:\s*[,،]\s*\d+)*)\]")


def plain(text: str) -> str:
    """Markdown emphasis and citation marks removed: a file shows text, not syntax."""
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return " ".join(text.split())


def parse(markdown: str, keep_citations: bool = True) -> list[Block]:
    blocks: list[Block] = []
    table: list[list[str]] = []

    def flush_table() -> None:
        if table:
            blocks.append(Block("table", rows=[row[:] for row in table]))
            table.clear()

    for raw in markdown.splitlines():
        line = raw.rstrip()
        if _TABLE_RULE.match(line) and table:
            continue
        row = _TABLE_ROW.match(line)
        if row:
            table.append([plain(c) for c in row.group(1).split("|")])
            continue
        flush_table()
        if not line.strip():
            continue
        if not keep_citations:
            line = _CITATION.sub("", line)
        heading = re.match(r"^\s*(#{1,4})\s+(.*)$", line)
        bullet = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)$", line)
        if heading:
            blocks.append(Block("heading", plain(heading.group(2)), level=len(heading.group(1))))
        elif bullet:
            blocks.append(Block("bullet", plain(bullet.group(1))))
        else:
            blocks.append(Block("paragraph", plain(line)))
    flush_table()
    return blocks


_NUMBER = re.compile(r"\d[\d,.]*")
_ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")


def number_supported(cell: str, evidence: str) -> bool:
    """Every number in the cell appears in the evidence, written either way."""
    numbers = _NUMBER.findall(cell.translate(_ARABIC_DIGITS))
    if not numbers:
        return True
    haystack = evidence.translate(_ARABIC_DIGITS)
    flat = haystack.replace(",", "")
    for number in numbers:
        clean = number.strip(".,")
        if not clean:
            continue
        if clean not in haystack and clean.replace(",", "") not in flat:
            return False
    return True


# -- the files -------------------------------------------------------------

FONT_DIR = Path(r"C:\Windows\Fonts")
_RTL = re.compile(r"[؀-ۿ]")


def _is_arabic(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and sum(1 for c in letters if _RTL.match(c)) / len(letters) >= 0.3


def render_pdf(content: Content, path: Path) -> None:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    rtl = _is_arabic(content.title + " ".join(b.text for b in content.blocks[:5]))
    align = "R" if rtl else "L"
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    regular, bold = FONT_DIR / "tahoma.ttf", FONT_DIR / "tahomabd.ttf"
    if not regular.exists():
        regular = bold = FONT_DIR / "arial.ttf"
    pdf.add_font("body", "", str(regular))
    pdf.add_font("body", "B", str(bold if bold.exists() else regular))
    pdf.set_text_shaping(use_shaping_engine=True, direction="rtl" if rtl else "ltr",
                         script="arab" if rtl else None, language="ar" if rtl else None)
    pdf.set_footer = None
    pdf.add_page()
    pdf.set_font("body", "B", 16)
    pdf.multi_cell(0, 9, content.title, align=align, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("body", "", 9)
    pdf.set_text_color(110, 110, 110)
    stamp = f"{datetime.now():%Y-%m-%d %H:%M}" + (f" — {content.author}" if content.author else "")
    pdf.multi_cell(0, 6, stamp, align=align, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)

    for block in content.blocks:
        if block.kind == "heading":
            pdf.ln(2)
            pdf.set_font("body", "B", 14 if block.level <= 2 else 12)
            pdf.multi_cell(0, 8, block.text, align=align, new_x="LMARGIN", new_y="NEXT")
        elif block.kind == "paragraph":
            pdf.set_font("body", "", 11)
            pdf.multi_cell(0, 7, block.text, align=align, new_x="LMARGIN", new_y="NEXT")
        elif block.kind == "bullet":
            pdf.set_font("body", "", 11)
            pdf.multi_cell(0, 7, (f"{block.text} •" if False else f"• {block.text}"), align=align,
                           new_x="LMARGIN", new_y="NEXT")
        elif block.kind == "table" and block.rows:
            pdf.ln(1)
            pdf.set_font("body", "", 9)
            width = max(len(r) for r in block.rows)
            rows = [r + [""] * (width - len(r)) for r in block.rows]
            # Right-to-left: the first column belongs on the right.
            rows = [list(reversed(r)) for r in rows] if rtl else rows
            with pdf.table(text_align="RIGHT" if rtl else "LEFT", line_height=6,
                           headings_style=FontFace(emphasis="BOLD", fill_color=(235, 238, 245))) as table:
                for row in rows:
                    cells = table.row()
                    for cell in row:
                        cells.cell(cell)
            pdf.ln(2)

    if content.note:
        pdf.ln(2)
        pdf.set_font("body", "", 9)
        pdf.set_text_color(150, 95, 0)
        pdf.multi_cell(0, 6, content.note, align=align, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
    if content.sources:
        pdf.ln(3)
        pdf.set_font("body", "B", 11)
        pdf.multi_cell(0, 7, "المصادر" if rtl else "Sources", align=align, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("body", "", 9)
        for index, source in enumerate(content.sources, 1):
            pdf.multi_cell(0, 6, f"[{index}] {source}", align=align, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def render_docx(content: Content, path: Path) -> None:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor

    rtl = _is_arabic(content.title + " ".join(b.text for b in content.blocks[:5]))
    document = Docx()
    style = document.styles["Normal"]
    style.font.name = "Tahoma"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:cs"), "Tahoma")

    def direct(paragraph) -> None:
        if not rtl:
            return
        properties = paragraph._p.get_or_add_pPr()
        properties.append(OxmlElement("w:bidi"))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in paragraph.runs:
            run_properties = run._r.get_or_add_rPr()
            run_properties.append(OxmlElement("w:rtl"))

    direct(document.add_heading(content.title, level=0))
    stamp = document.add_paragraph(f"{datetime.now():%Y-%m-%d %H:%M}" + (f" — {content.author}" if content.author else ""))
    stamp.runs[0].font.color.rgb = RGBColor(110, 110, 110)
    direct(stamp)
    for block in content.blocks:
        if block.kind == "heading":
            direct(document.add_heading(block.text, level=min(3, max(1, block.level))))
        elif block.kind == "paragraph":
            direct(document.add_paragraph(block.text))
        elif block.kind == "bullet":
            direct(document.add_paragraph(block.text, style="List Bullet"))
        elif block.kind == "table" and block.rows:
            width = max(len(r) for r in block.rows)
            table = document.add_table(rows=0, cols=width)
            table.style = "Light Grid Accent 1"
            if rtl:
                table_properties = table._tbl.tblPr
                bidi = OxmlElement("w:bidiVisual")
                table_properties.append(bidi)
            for index, row in enumerate(block.rows):
                cells = table.add_row().cells
                for column, value in enumerate(row + [""] * (width - len(row))):
                    cells[column].text = value
                    for paragraph in cells[column].paragraphs:
                        direct(paragraph)
                        if index == 0:
                            for run in paragraph.runs:
                                run.bold = True
    if content.note:
        note = document.add_paragraph(content.note)
        note.runs[0].font.color.rgb = RGBColor(150, 95, 0)
        direct(note)
    if content.sources:
        direct(document.add_heading("المصادر" if rtl else "Sources", level=2))
        for index, source in enumerate(content.sources, 1):
            direct(document.add_paragraph(f"[{index}] {source}"))
    document.save(str(path))


def render_xlsx(content: Content, path: Path) -> int:
    """The tables as sheets — or, with none, the content as rows. Returns how many
    numeric cells were not found in the evidence and were marked."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rtl = _is_arabic(content.title + " ".join(b.text for b in content.blocks[:5]))
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    unsure_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="C9CED6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    unsupported = 0

    tables = [b.rows for b in content.blocks if b.kind == "table" and b.rows]
    if not tables:
        rows = [["البند" if rtl else "Item", "التفاصيل" if rtl else "Details"]]
        for block in content.blocks:
            if block.kind in ("paragraph", "bullet", "heading"):
                label, sep, value = block.text.partition(":")
                rows.append([label.strip(), value.strip()] if sep and len(label) <= 60 else [block.text, ""])
        tables = [rows]

    def number(value: str):
        compact = value.translate(_ARABIC_DIGITS).replace(",", "").strip()
        try:
            return float(compact) if re.fullmatch(r"-?\d+(\.\d+)?", compact) else None
        except ValueError:
            return None

    for index, rows in enumerate(tables, 1):
        sheet = workbook.create_sheet(title=(f"جدول {index}" if rtl else f"Table {index}")[:31])
        sheet.sheet_view.rightToLeft = rtl
        sheet.append([content.title])
        sheet["A1"].font = Font(bold=True, size=14)
        sheet.append([f"{datetime.now():%Y-%m-%d %H:%M}"])
        sheet["A2"].font = Font(color="808080", size=9)
        sheet.append([])
        width = max(len(r) for r in rows)
        start = sheet.max_row + 1
        for row_index, row in enumerate(rows):
            values = row + [""] * (width - len(row))
            sheet.append([number(v) if row_index and number(v) is not None else v for v in values])
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row=sheet.max_row, column=column)
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical="top",
                                           horizontal="right" if rtl else "left")
                if row_index == 0:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                elif content.evidence and not number_supported(value, content.evidence):
                    cell.fill = unsure_fill
                    cell.comment = Comment("لم يُعثر على هذا الرقم في المصادر — راجعه." if rtl
                                           else "Not found in the sources — check it.", "RAG")
                    unsupported += 1
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.##"
        sheet.freeze_panes = sheet.cell(row=start + 1, column=1)
        for column in range(1, width + 1):
            longest = max((len(str(r[column - 1])) if column - 1 < len(r) else 0) for r in rows)
            sheet.column_dimensions[get_column_letter(column)].width = min(60, max(10, longest + 2))

    if content.sources or content.note:
        sheet = workbook.create_sheet(title="المصادر" if rtl else "Sources")
        sheet.sheet_view.rightToLeft = rtl
        if content.note:
            sheet.append([content.note])
        for index, source in enumerate(content.sources, 1):
            sheet.append([f"[{index}]", source])
        sheet.column_dimensions["B"].width = 80
    workbook.save(str(path))
    return unsupported


# -- keeping them ----------------------------------------------------------

@dataclass
class StoredFile:
    id: str
    owner: str
    name: str
    format: str
    size: int
    created: str

    @property
    def media_type(self) -> str:
        return MEDIA_TYPES[self.format]


class FileStore:
    """One folder per person. A file is served only to whoever it was made for."""

    def __init__(self, root: Path) -> None:
        self.root = root

    def _folder(self, owner: str) -> Path:
        safe = re.sub(r"[^\w-]", "_", owner)
        return self.root / safe

    def path_for(self, owner: str, file_id: str, fmt: str) -> Path:
        folder = self._folder(owner)
        folder.mkdir(parents=True, exist_ok=True)
        return folder / f"{file_id}.{fmt}"

    def new_id(self) -> str:
        return uuid.uuid4().hex

    def record(self, owner: str, file_id: str, name: str, fmt: str) -> StoredFile:
        path = self.path_for(owner, file_id, fmt)
        stored = StoredFile(file_id, owner, name, fmt, path.stat().st_size, datetime.now(UTC).isoformat())
        (path.with_suffix(".json")).write_text(json.dumps(stored.__dict__, ensure_ascii=False), encoding="utf-8")
        return stored

    def open(self, owner: str, file_id: str) -> tuple[StoredFile, Path]:
        if not re.fullmatch(r"[0-9a-f]{32}", file_id):
            raise DocumentNotFoundError("الملف غير موجود.")
        meta = self._folder(owner) / f"{file_id}.json"
        if not meta.exists():
            raise DocumentNotFoundError("الملف غير موجود أو انتهت مدته.")
        stored = StoredFile(**json.loads(meta.read_text(encoding="utf-8")))
        path = self.path_for(owner, file_id, stored.format)
        if stored.owner != owner or not path.exists():
            raise DocumentNotFoundError("الملف غير موجود أو انتهت مدته.")
        return stored, path

    def remove_older_than(self, days: int = KEEP_DAYS) -> int:
        if not self.root.exists():
            return 0
        limit = datetime.now().timestamp() - timedelta(days=days).total_seconds()
        removed = 0
        for path in self.root.glob("*/*"):
            if path.is_file() and path.stat().st_mtime < limit:
                path.unlink(missing_ok=True)
                removed += path.suffix != ".json"
        return removed
