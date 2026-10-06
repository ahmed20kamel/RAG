"""Accuracy on the company's own files, measured against answers a person approved.

Two steps, because the right answers are not the system's to decide:

1. ``--draft`` answers every question in ``data/eval/company_questions.json`` and writes a
   review workbook: the question, the system's answer, its sources, and empty columns for
   the reviewer — correct / wrong / incomplete, the right answer, and the key facts (the
   figures, names and dates) any correct answer must contain.

2. ``--score`` reads the reviewed workbook back, answers the questions again with the
   system as it is now, and counts how many answers contain every key fact the reviewer
   wrote. Run after every change to the pipeline: a change that lowers the score is a
   regression, whatever it was meant to improve.

The questions and the workbook stay in ``data/`` — case material, never committed.

Run: python scripts/company_eval.py --draft
     python scripts/company_eval.py --score
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

QUESTIONS = ROOT / "data" / "eval" / "company_questions.json"
WORKBOOK = ROOT / "data" / "eval" / "مراجعة_أسئلة_الدقة.xlsx"
RESULTS = ROOT / "data" / "eval" / "results"
HEADERS = ["الرقم", "الفئة", "السؤال", "إجابة النظام", "المصادر", "التقييم (صح / خطأ / ناقصة)",
           "الإجابة الصحيحة أو الملاحظة", "الحقائق الأساسية (افصل بينها بـ ؛)"]

_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def _answer_all(questions: list[dict]) -> list[dict]:
    from app.config import get_settings
    from app.container import Container
    from app.schemas.chat import ChatRequest

    service = Container(get_settings()).rag_service
    rows = []
    for index, item in enumerate(questions, 1):
        started = time.perf_counter()
        try:
            response = service.answer(ChatRequest(question=item["question"], fresh=True), user=None)
            answer, sources = response.answer, sorted({
                s.filename + (f" — {s.locator}" if s.locator else "") for s in response.sources
            })
        except Exception as exc:  # noqa: BLE001 - one failure is a row, not the end of the run
            answer, sources = f"خطأ: {exc}", []
        seconds = time.perf_counter() - started
        print(f"[{index}/{len(questions)}] {item['id']} ({seconds:.0f}s) {item['question'][:60]}", flush=True)
        rows.append({**item, "answer": answer, "sources": sources, "seconds": round(seconds)})
    return rows


def draft() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    questions = json.loads(QUESTIONS.read_text(encoding="utf-8"))
    rows = _answer_all(questions)
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / f"draft-{datetime.now():%Y%m%d-%H%M}.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")

    book = Workbook()
    sheet = book.active
    sheet.title = "المراجعة"
    sheet.sheet_view.rightToLeft = True
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        sheet.append([row["id"], row["category"], row["question"], row["answer"], "\n".join(row["sources"]), "", "", ""])
    verdict = DataValidation(type="list", formula1='"صح,خطأ,ناقصة"', allow_blank=True)
    sheet.add_data_validation(verdict)
    verdict.add(f"F2:F{len(rows) + 1}")
    for column, width in zip("ABCDEFGH", (7, 14, 38, 80, 34, 16, 50, 40)):
        sheet.column_dimensions[column].width = width
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    book.save(WORKBOOK)
    print(f"\nReview workbook: {WORKBOOK}")


def _facts(cell: str) -> list[str]:
    return [f.strip() for f in re.split(r"[؛;\n]", cell or "") if f.strip()]


def _contains(answer: str, fact: str) -> bool:
    def norm(text: str) -> str:
        text = text.translate(_DIGITS).replace("،", ",")
        text = re.sub(r"[ًٌٍَُِّْـ]", "", text)
        text = re.sub(r"[إأآ]", "ا", text).replace("ة", "ه").replace("ى", "ي")
        return " ".join(text.split())
    a, f = norm(answer), norm(fact)
    return f in a or f.replace(",", "") in a.replace(",", "")


def score() -> None:
    from openpyxl import load_workbook

    sheet = load_workbook(WORKBOOK).active
    reviewed = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        qid, category, question, _answer, _sources, verdict, _note, facts = (list(row) + [None] * 8)[:8]
        if question and _facts(facts or ""):
            reviewed.append({"id": qid, "category": category, "question": question,
                             "verdict": verdict, "facts": _facts(facts)})
    if not reviewed:
        print("No question has its key facts filled in yet — review the workbook first.")
        return
    rows = _answer_all(reviewed)
    passed = 0
    for row in rows:
        missing = [f for f in row["facts"] if not _contains(row["answer"], f)]
        row["missing"] = missing
        passed += not missing
        print(f"  {'✅' if not missing else '❌'} {row['id']} {row['question'][:50]}" + (f"  — ناقص: {missing}" if missing else ""))
    accuracy = passed / len(rows)
    RESULTS.mkdir(parents=True, exist_ok=True)
    (RESULTS / f"score-{datetime.now():%Y%m%d-%H%M}.json").write_text(
        json.dumps({"accuracy": accuracy, "passed": passed, "total": len(rows), "rows": rows},
                   ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\nAccuracy: {passed}/{len(rows)} = {accuracy:.0%}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--draft", action="store_true")
    group.add_argument("--score", action="store_true")
    args = parser.parse_args()
    draft() if args.draft else score()
