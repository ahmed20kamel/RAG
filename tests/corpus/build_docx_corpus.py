"""Word documents built to have the shapes the real ones turned out to have.

Two of these matter more than the rest. `styled_headings.docx` is the textbook case that
any parser handles. `bold_headings.docx` is what the company actually writes: every
paragraph styled `Normal`, the headings bold and numbered, and one bold line in the middle
that is emphasis rather than a heading — the exact thing that makes a naive "bold means
heading" rule produce a shredded document.

    python tests/corpus/build_docx_corpus.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
sys.stdout.reconfigure(encoding="utf-8")


def styled_headings() -> Document:
    """Real heading styles, a nested tree, and a table under a known heading."""
    document = Document()
    document.core_properties.title = "عقد المقاولة"
    document.core_properties.author = "الإدارة القانونية"
    document.core_properties.category = "عقود"

    document.add_heading("عقد المقاولة", level=1)
    document.add_paragraph("هذا العقد مبرم بين الطرفين المذكورين أدناه.")

    document.add_heading("شروط الدفع", level=2)
    document.add_paragraph(
        "تُصرف الدفعات على أساس شهري بعد اعتماد المستخلص من الاستشاري."
    )
    table = document.add_table(rows=4, cols=3)
    data = [
        ["الدفعة", "النسبة", "المبلغ"],
        ["الدفعة الأولى", "30%", "804000"],
        ["الدفعة الثانية", "45%", "1206000"],
        ["الدفعة النهائية", "25%", "670000"],
    ]
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            cell.text = value

    document.add_heading("غرامة التأخير", level=2)
    document.add_paragraph(
        "تُحتسب غرامة التأخير بواقع 1,985.19 درهم عن كل يوم تأخير، "
        "بحد أقصى 10% من قيمة العقد."
    )

    document.add_heading("إنهاء العقد", level=2)
    document.add_paragraph("لصاحب العمل إنهاء العقد بإخطار خطي مدته أربعة عشر يومًا.")
    return document


def bold_headings() -> Document:
    """No heading styles at all — numbered bold lines, plus one bold line that is not one."""
    document = Document()

    def bold(text: str) -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = True

    bold("التقرير الكامل — العمليات والفرق")
    bold("1. الهدف من العمل")
    document.add_paragraph("كان الهدف ألا يبقى النظام مجرد أداة تسجيل بسيطة.")
    # Bold, short, unpunctuated — and emphasis, not a heading. A parser that takes every
    # bold line as a heading breaks the section here and buries the rest of section 1.
    bold("Team = Supervisor + Workers")
    document.add_paragraph("بل يتحول إلى نظام متكامل يصلح لأي شركة خدمات.")

    bold("2. المشكلة القديمة")
    document.add_paragraph("كانت البيانات موزعة على جداول غير مترابطة.")
    document.add_paragraph("ولم يكن هناك مصدر واحد للحقيقة في تكاليف المشاريع.")

    bold("3. الحل المعتمد")
    document.add_paragraph("اعتُمد نموذج موحّد يربط الفرق بالمشاريع ومراكز التكلفة.")
    document.add_paragraph("وصار رقم المرجع الداخلي هو OPS-2026-4471.")
    return document


def mixed_languages() -> Document:
    """Arabic and English in the same document, and in the same table."""
    document = Document()
    document.add_heading("Scope of Work — نطاق الأعمال", level=1)
    document.add_paragraph(
        "The contractor shall complete all structural works — يلتزم المقاول بإنجاز "
        "الأعمال الإنشائية كاملة."
    )

    document.add_heading("Deliverables — المخرجات", level=2)
    table = document.add_table(rows=4, cols=3)
    data = [
        ["Item", "الوصف", "Duration"],
        ["Excavation", "أعمال الحفر", "21 days"],
        ["Foundation", "الأساسات", "45 days"],
        ["Superstructure", "الهيكل العلوي", "120 days"],
    ]
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            cell.text = value
    return document


def tables_only() -> Document:
    """Two tables under one heading — the case where naming a table would be a guess."""
    document = Document()
    document.add_heading("الملاحق", level=1)
    for index, rows in enumerate(
        [
            [["الرقم", "البند"], ["1", "ملحق أول"], ["2", "ملحق ثانٍ"]],
            [["Code", "Value"], ["A1", "100"], ["A2", "250"]],
        ],
        start=1,
    ):
        document.add_paragraph(f"الملحق رقم {index}")
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row, values in zip(table.rows, rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
    return document


def interleaved() -> Document:
    """A table in the middle, to catch a parser that reads all paragraphs then all tables."""
    document = Document()
    document.add_heading("القسم الأول", level=1)
    document.add_paragraph("نص القسم الأول قبل الجدول.")

    document.add_heading("القسم الثاني", level=1)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "المفتاح"
    table.rows[0].cells[1].text = "القيمة"
    table.rows[1].cells[0].text = "رقم الإذن"
    table.rows[1].cells[1].text = "PTW-9931"
    document.add_paragraph("نص يلي الجدول في القسم الثاني.")

    document.add_heading("القسم الثالث", level=1)
    document.add_paragraph("نص القسم الثالث بعد كل ما سبق.")
    return document


def long_document() -> Document:
    """Long enough that flat extraction would be useless."""
    document = Document()
    for chapter in range(1, 21):
        document.add_heading(f"الفصل {chapter}", level=1)
        for paragraph in range(4):
            document.add_paragraph(
                f"الفقرة {paragraph + 1} من الفصل {chapter}. "
                f"القيمة المرجعية لهذا الفصل هي {chapter * 1000}. " * 3
            )
    return document


def empty_document() -> Document:
    return Document()


BUILDERS = {
    "styled_headings.docx": styled_headings,
    "bold_headings.docx": bold_headings,
    "mixed_languages.docx": mixed_languages,
    "tables_only.docx": tables_only,
    "interleaved.docx": interleaved,
    "long_document.docx": long_document,
    "empty.docx": empty_document,
}


def build() -> None:
    HERE.mkdir(parents=True, exist_ok=True)
    for name, builder in BUILDERS.items():
        document = builder()
        document.save(HERE / name)
        print(f"  {name:26} {(HERE / name).stat().st_size:>8,} bytes")

    (HERE / "corrupt.docx").write_bytes(b"PK\x03\x04" + b"\x00" * 200)
    # A workbook the end-to-end suite does not also upload as itself: using
    # financial_summary.xlsx here made this a byte-identical duplicate of a document
    # already indexed, and the upload was correctly rejected as a duplicate before
    # the type check it was meant to exercise ever ran.
    (HERE / "renamed_xlsx.docx").write_bytes((HERE / "large_ledger.xlsx").read_bytes())
    print("  corrupt.docx / renamed_xlsx.docx  (rejection cases)")


if __name__ == "__main__":
    print("building the DOCX corpus\n")
    build()
    print(f"\nwritten to {HERE}")
