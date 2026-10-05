"""Files asked for in the chat: recognised, built correctly, and kept to their owner.

A request for a file is told from a question that merely names one ("ما في ملف
العقد.pdf؟"); the question inside the request is recovered intact; each format opens and
holds the content — Arabic shaped in the PDF, a right-to-left sheet in Excel; a number the
sources do not contain is highlighted rather than presented as quoted; and a file is
served to the person it was made for and to no one else.

Offline. Run: python tests/test_file_export.py
"""

from __future__ import annotations

import io
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)

from app.exceptions import DocumentNotFoundError  # noqa: E402
from app.services.file_export import (  # noqa: E402
    Content, FileStore, detect, number_supported, parse, render_docx, render_pdf, render_xlsx,
)

FAILURES: list[str] = []


def check(condition: bool, label: str, detail: str = "") -> None:
    print(f"[{'PASS' if condition else 'FAIL'}] {label}")
    if not condition:
        if detail:
            print(f"        {detail}")
        FAILURES.append(label)


ANSWER = """## أسعار الدهان

سعر الدهان الداخلي في البنود المضافة **45 درهم/م²** [1].

| البند | السعر | المصدر |
|---|---|---|
| دهان داخلي | 45 | أمر التغيير |
| دهان خارجي | 62 | أمر التغيير |

- نسبة الزيادة حسب جوتن 7% [2]
"""
EVIDENCE = "البنود المضافة: دهان داخلي 45 درهم للمتر المربع. نسبة الزيادة 7% حسب جوتن."


def main() -> int:
    print("=== 1. what counts as a file request ===")
    cases = {
        "حول دا لـ PDF": ("pdf", "last", ""),
        "هاتلي ملخص المحادثة PDF": ("pdf", "conversation", ""),
        "export the conversation to pdf": ("pdf", "conversation", ""),
        "ابعتلي الإجابة word": ("docx", "last", ""),
        "give me this as an excel file": ("xlsx", "last", ""),
        "اعمل ملف اكسيل فيه مقارنة بين عرض جوتن وعرض ناشيونال": ("xlsx", "question", "مقارنة بين عرض جوتن وعرض ناشيونال؟"),
        "اعملي بي دي اف بأسعار الدهان الداخلي في أمر التغيير": ("pdf", "question", "أسعار الدهان الداخلي في أمر التغيير؟"),
    }
    for text, expected in cases.items():
        found = detect(text)
        got = (found.format, found.subject, found.question) if found else None
        check(got == expected, f"«{text}»", f"{got} != {expected}")
    for text in ("ما في ملف العقد.pdf؟", "هل يوجد ملف pdf عن الغرامات؟", "ما قيمة العقد؟"):
        check(detect(text) is None, f"«{text}» is a question, not a request for a file")

    print("\n=== 2. the content ===")
    blocks = parse(ANSWER)
    kinds = [b.kind for b in blocks]
    check(kinds == ["heading", "paragraph", "table", "bullet"], "headings, text, tables and lists are told apart", str(kinds))
    check("**" not in blocks[1].text, "emphasis marks do not reach the file")
    check(blocks[2].rows[0] == ["البند", "السعر", "المصدر"] and len(blocks[2].rows) == 3, "the table keeps its header and rows")
    check(number_supported("45", EVIDENCE) and number_supported("7%", EVIDENCE), "numbers in the sources are supported")
    check(not number_supported("62", EVIDENCE), "a number the sources do not contain is not")

    print("\n=== 3. the files ===")
    work = Path(tempfile.mkdtemp(prefix="rag-files-"))
    try:
        content = Content(title="أسعار الدهان — أمر التغيير", blocks=blocks, sources=["أمر التغيير.pdf — صفحة 2"],
                          evidence=EVIDENCE, author="tester")
        pdf, docx, xlsx = work / "a.pdf", work / "a.docx", work / "a.xlsx"
        render_pdf(content, pdf)
        render_docx(content, docx)
        unsupported = render_xlsx(content, xlsx)

        import pymupdf

        text = pymupdf.open(str(pdf))[0].get_text()
        check("دهان" in text and "45" in text, "the PDF opens and holds the Arabic text and figures", text[:80])

        from docx import Document

        document = Document(str(docx))
        paragraphs = " ".join(p.text for p in document.paragraphs)
        table_text = [c.text for t in document.tables for r in t.rows for c in r.cells]
        check("أسعار الدهان" in paragraphs and "دهان داخلي" in table_text,
              "the Word file holds the text and the table")
        check(any("ال يافور" in p.text for p in document.sections[0].header.paragraphs),
              "every page carries the company's name in its header")
        check(document.paragraphs[0]._p.pPr is not None and document.paragraphs[0]._p.pPr.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None,
              "Word paragraphs are right-to-left")

        from openpyxl import load_workbook

        book = load_workbook(str(xlsx))
        sheet = book.worksheets[0]
        values = [[c.value for c in row] for row in sheet.iter_rows()]
        flat = [v for row in values for v in row if v is not None]
        check("دهان داخلي" in flat and 45 in flat, "the sheet holds the rows, numbers as numbers", str(flat[:8]))
        check(sheet.sheet_view.rightToLeft, "the sheet reads right to left")
        check(unsupported == 1, "one number without support is counted", str(unsupported))
        marked = [c for row in sheet.iter_rows() for c in row if c.comment is not None]
        check(len(marked) == 1 and marked[0].value == 62, "…and highlighted with a note, not left looking quoted")
        check(book.sheetnames[-1] == "المصادر", "the sources travel with the file")

        print("\n=== 4. whose file it is ===")
        store = FileStore(work / "exports")
        file_id = store.new_id()
        path = store.path_for("alice", file_id, "pdf")
        path.write_bytes(pdf.read_bytes())
        stored = store.record("alice", file_id, "أسعار.pdf", "pdf")
        check(store.open("alice", file_id)[0].name == "أسعار.pdf", "the owner gets their file")
        try:
            store.open("bob", file_id)
            check(False, "someone else asking for it gets nothing")
        except DocumentNotFoundError:
            check(True, "someone else asking for it gets nothing")
        try:
            store.open("alice", "../../etc")
            check(False, "a path in place of an id is refused")
        except DocumentNotFoundError:
            check(True, "a path in place of an id is refused")
        check(stored.size > 0, "the size is recorded")
    finally:
        shutil.rmtree(work, ignore_errors=True)

    print()
    if FAILURES:
        print(f"{len(FAILURES)} FAILURE(S)")
        return 1
    print("all checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
